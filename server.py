from flask import Flask, send_from_directory, request
from pptx import Presentation
from PyPDF2 import PdfReader
import markdown
import os
import base64
import re

app = Flask(__name__)

FILES_FOLDER = "files"


def search_files(keyword):

    keyword = keyword.lower()

    results = []

    for filename in os.listdir(FILES_FOLDER):

        path = os.path.join(
            FILES_FOLDER,
            filename
        )

        text = ""

        try:

            # Markdown
            if filename.lower().endswith(".md"):

                with open(
                    path,
                    "r",
                    encoding="utf-8"
                ) as f:

                    text = f.read()

            # HTML
            elif filename.lower().endswith(".html"):

                with open(
                    path,
                    "r",
                    encoding="utf-8",
                    errors="ignore"
                ) as f:

                    text = f.read()

            # PDF
            elif filename.lower().endswith(".pdf"):

                pdf = PdfReader(path)

                for page_num, page in enumerate(pdf.pages):

                    extracted = page.extract_text()

                    if extracted:

                        text += (
                            f"\n[Page {page_num+1}]\n"
                            + extracted
                        )

            # PPT / PPTX
            elif (
                filename.lower().endswith(".ppt")
                or
                filename.lower().endswith(".pptx")
            ):

                prs = Presentation(path)

                for slide_num, slide in enumerate(prs.slides):

                    text += f"\n[Slide {slide_num+1}]\n"

                    for shape in slide.shapes:

                        if hasattr(
                            shape,
                            "text"
                        ):

                            text += (
                                shape.text
                                + "\n"
                            )

        except:

            continue

        paragraphs = re.split(
            r"\n+",
            text
        )

        for para in paragraphs:

            if keyword in para.lower():

                if filename.lower().endswith(".md"):

                    link = f"/view/{filename}"

                elif filename.lower().endswith(".pdf"):

                    link = f"/pdf/{filename}"

                elif filename.lower().endswith(".html"):

                    link = f"/html/{filename}"

                else:

                    link = f"/ppt/{filename}"

                results.append({

                    "file": filename,

                    "snippet": para[:350],

                    "link": link

                })

    return results


@app.route("/")
def home():

    files = os.listdir(
        FILES_FOLDER
    )

    html = """

<html>

<style>

body{
background:#0d1117;
font-family:Arial;
padding:40px;
color:white;
}

.file{
display:block;
padding:12px;
margin:10px 0;
background:#161b22;
border-radius:10px;
text-decoration:none;
color:#58a6ff;
}

.searchLink{
color:#58a6ff;
}

</style>

<body>

<h1>StudyCircle</h1>

<a class="searchLink"
href="/topics">

Study Topics

</a>

<br><br>

<a class="searchLink"
href="/search">

Search Content

</a>

<br><br>

"""

    for file in files:

        lower = file.lower()

        if lower.endswith(".md"):

            link = f"/view/{file}"

        elif lower.endswith(".pdf"):

            link = f"/pdf/{file}"

        elif lower.endswith(".html"):

            link = f"/html/{file}"

        elif (
            lower.endswith(".ppt")
            or
            lower.endswith(".pptx")
        ):

            link = f"/ppt/{file}"

        else:

            continue

        html += f"""

<a
class="file"
href="{link}">

{file}

</a>

"""

    html += "</body></html>"

    return html


@app.route("/topics")
def topics():

    return """

<html>

<body
style="
background:#0d1117;
padding:40px;
color:white;
font-family:Arial;
">

<a href="/">← Back</a>

<h1>Study Topics</h1>

<input id="topicInput">

<button onclick="addTopic()">

Add

</button>

<div id="topics"></div>

<script>

const area =
document.getElementById(
"topics"
);

function load(){

area.innerHTML="";

const topics =
JSON.parse(

localStorage.getItem(
"study_topics"
)

||

"[]"

);

topics.forEach(

(t,i)=>{

area.innerHTML += `

<div style="
padding:10px;
margin:10px 0;
background:#161b22;
">

<input

type="checkbox"

${t.done?'checked':''}

onchange="toggle(${i})"

>

<span style="
${t.done?
'opacity:.5;text-decoration:line-through'
:''}
">

${t.name}

</span>

</div>

`;

});

}

function addTopic(){

const input =
document.getElementById(
"topicInput"
);

if(!input.value)
return;

const topics =
JSON.parse(

localStorage.getItem(
"study_topics"
)

||

"[]"

);

topics.push({

name:input.value,

done:false

});

localStorage.setItem(

"study_topics",

JSON.stringify(
topics
)

);

input.value="";

load();

}

function toggle(i){

const topics =
JSON.parse(

localStorage.getItem(
"study_topics"
)

||

"[]"

);

topics[i].done =
!topics[i].done;

localStorage.setItem(

"study_topics",

JSON.stringify(
topics
)

);

load();

}

load();

</script>

</body>

</html>

"""


@app.route("/search")
def search():

    query = request.args.get(
        "q",
        ""
    )

    output = ""

    if query:

        results = search_files(
            query
        )

        for r in results:

            output += f"""

<a
class="card"
href="{r['link']}">

<h3>

{r['file']}

</h3>

<p>

{r['snippet']}

</p>

</a>

"""

    return f"""

<html>

<style>

body{{
background:#0d1117;
padding:40px;
font-family:Arial;
color:white;
}}

input{{
width:100%;
padding:12px;
background:#161b22;
border:none;
color:white;
}}

.card{{
display:block;
background:#161b22;
padding:15px;
margin:10px 0;
text-decoration:none;
color:white;
}}

</style>

<body>

<a href="/">← Back</a>

<h1>Search</h1>

<form>

<input
name="q"
value="{query}"
placeholder="keyword">

</form>

{output}

</body>

</html>

"""


@app.route("/view/<filename>")
def view(filename):

    with open(
        os.path.join(
            FILES_FOLDER,
            filename
        ),
        encoding="utf-8"
    ) as f:

        md = f.read()

    rendered = markdown.markdown(
        md,
        extensions=[
            "tables",
            "fenced_code"
        ]
    )

    return f"""

<body
style="
background:#0d1117;
padding:40px;
color:white;
">

<a href="/">← Back</a>

{rendered}

</body>

"""


@app.route("/html/<filename>")
def html_view(filename):

    with open(
        os.path.join(
            FILES_FOLDER,
            filename
        ),
        encoding="utf-8",
        errors="ignore"
    ) as f:

        html = f.read()

    return html


@app.route("/pdf/<filename>")
def pdf_view(filename):

    return f"""

<object
data="/pdf_file/{filename}"
type="application/pdf"
style="
width:100%;
height:100vh;
">

</object>

"""


@app.route("/pdf_file/<filename>")
def pdf_file(filename):

    return send_from_directory(
        FILES_FOLDER,
        filename
    )


@app.route("/ppt/<filename>")
def ppt_view(filename):

    prs = Presentation(
        os.path.join(
            FILES_FOLDER,
            filename
        )
    )

    content = ""

    for slide in prs.slides:

        for shape in slide.shapes:

            if hasattr(
                shape,
                "text"
            ):

                content += f"<p>{shape.text}</p>"

            if shape.shape_type == 13:

                try:

                    img = base64.b64encode(
                        shape.image.blob
                    ).decode()

                    ext = shape.image.ext

                    content += f'''

<img
src="data:image/{ext};base64,{img}"
style="max-width:100%">

'''

                except:

                    pass

    return f"""

<body
style="
background:#0d1117;
padding:40px;
color:white;
">

<a href="/">← Back</a>

{content}

</body>

"""


app.run(
host="0.0.0.0",
port=8000,
debug=True
)